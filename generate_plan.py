import json
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

with open('parsed_items.json', 'r', encoding='utf-8') as f:
    items = json.load(f)

with open('complete_mapping.json', 'r', encoding='utf-8') as f:
    mappings = json.load(f)

items_by_num = {it['num']: it for it in items}

records = []
for m in mappings:
    it = items_by_num[m['num']]
    feature_id_name = f"{it['id']} - {it['title']}"
    workstream = it['domain']
    # standardize workstream naming to match catalog where appropriate
    ws_map = {
        'HCM': 'HCM Core',
        'UK Payroll': 'UK Payroll (Gap)',
        'People Experience': 'People Experience (Help)',
        'Cross App': 'Cross Application',
        'Extend': 'Workday Extend'
    }
    std_workstream = ws_map.get(workstream, workstream)
    
    rec = {
        "Release Item ID / Feature Name": feature_id_name,
        "Workstream": std_workstream,
        "Impacted Workday BP(s)": m['impacted_bp'],
        "Targeted E2E Job ID & Name": m['job_id_name'],
        "Selection Priority (High/Med/Low)": m['priority'],
        "Risk & Impact Justification (\"Why it must be tested\")": m['justification'],
        "Impact Type": m['impact_type'],
        "Delivery Type": it['mandatory'],
        "Release Number": "2026R2",
        "Release Note URL": f"https://doc.workday.com/release-notes/en-us/hrcore/{it['id']}.html" if it.get('id') else ""
    }
    records.append(rec)

df = pd.DataFrame(records)

# Save to CSV
csv_path = 'Targeted_Workday_Release_Regression_Plan.csv'
df.to_csv(csv_path, index=False, encoding='utf-8-sig')
print(f"Saved {csv_path}")

# Save to Excel with openpyxl styling
excel_path = 'Targeted_Workday_Release_Regression_Plan.xlsx'
wb = openpyxl.Workbook()

# Sheet 1: Targeted Regression Plan
ws1 = wb.active
ws1.title = "Targeted Regression Plan"

headers = list(df.columns)
ws1.append(headers)

# Styling headers
header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

for col_idx, h in enumerate(headers, 1):
    cell = ws1.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Priority Fills
p_fills = {
    "High": PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
    "Med": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
    "Low": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
}
p_fonts = {
    "High": Font(name="Calibri", size=10, bold=True, color="C00000"),
    "Med": Font(name="Calibri", size=10, bold=True, color="B28900"),
    "Low": Font(name="Calibri", size=10, bold=True, color="375623")
}

gap_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
gap_font = Font(name="Calibri", size=10, italic=True, color="7F7F7F")

for row_idx, rec in enumerate(records, 2):
    for col_idx, h in enumerate(headers, 1):
        val = rec[h]
        cell = ws1.cell(row=row_idx, column=col_idx, value=val)
        cell.font = Font(name="Calibri", size=10)
        cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        
        # Color coding for priority
        if h == "Selection Priority (High/Med/Low)":
            p = str(val).strip()
            if p in p_fills:
                cell.fill = p_fills[p]
                cell.font = p_fonts[p]
            cell.alignment = Alignment(horizontal="center", vertical="top")
        
        # Highlight gaps
        if h == "Targeted E2E Job ID & Name" and "Gap Identified" in str(val):
            cell.fill = gap_fill
            cell.font = gap_font

# Adjust column widths
col_widths = {
    1: 38,  # Release Item
    2: 18,  # Workstream
    3: 32,  # Impacted BP
    4: 38,  # Targeted Job
    5: 18,  # Priority
    6: 55,  # Justification
    7: 22,  # Impact Type
    8: 16,  # Delivery Type
    9: 14,  # Release
    10: 30  # URL
}
for col_idx, width in col_widths.items():
    ws1.column_dimensions[get_column_letter(col_idx)].width = width

ws1.freeze_panes = "A2"
ws1.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(records)+1}"

# Sheet 2: Summary Dashboard & Metrics
ws2 = wb.create_sheet(title="Executive Summary")
ws2.views.sheetView[0].showGridLines = True

ws2['B2'] = "WORKDAY RELEASE IMPACT & REGRESSION SELECTION REPORT"
ws2['B2'].font = Font(name="Calibri", size=16, bold=True, color="1F497D")

ws2['B4'] = "Metric"
ws2['C4'] = "Value"
ws2['D4'] = "Context / Description"
for col_letter in ['B', 'C', 'D']:
    cell = ws2[f"{col_letter}4"]
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

with open('catalog_jobs.json', 'r', encoding='utf-8') as f:
    catalog_jobs = json.load(f)

unique_jobs = set(m['job_id_name'] for m in mappings if m['job_id_name'] != "No existing E2E Job found - Gap Identified")
total_items = len(records)
direct_impact_count = sum(1 for m in mappings if m['job_id_name'] != "No existing E2E Job found - Gap Identified")
gap_count = sum(1 for m in mappings if m['job_id_name'] == "No existing E2E Job found - Gap Identified")
total_catalog_jobs = len(catalog_jobs)
selected_jobs_count = len(unique_jobs)
reduction_pct = (1.0 - (selected_jobs_count / total_catalog_jobs)) * 100.0

high_priority_count = sum(1 for m in mappings if m['priority'] == 'High')
med_priority_count = sum(1 for m in mappings if m['priority'] == 'Med')
low_priority_count = sum(1 for m in mappings if m['priority'] == 'Low')

metrics = [
    ("Total Release Items Analyzed", total_items, "Total What's New items processed from release notes"),
    ("Release Items with Direct / Downstream E2E Test Impact", direct_impact_count, "Features mapped to existing regression suite"),
    ("Identified Coverage Gaps (Missing E2E Tests)", gap_count, "Release features lacking test coverage in catalog"),
    ("Master E2E Job Catalog Size", total_catalog_jobs, "Total test inventory available in Master Repository"),
    ("Targeted E2E Regression Jobs Selected", selected_jobs_count, "Unique tests scheduled for this release cycle"),
    ("Regression Suite Reduction Percentage", f"{reduction_pct:.1f}%", "Execution footprint reduction vs running entire catalog"),
    ("High Priority Selected Items", high_priority_count, "Critical business transactions and regulatory changes"),
    ("Medium Priority Selected Items", med_priority_count, "Configuration-dependent UI and calculation updates"),
    ("Low Priority Selected Items", low_priority_count, "Cosmetic or low-blast-radius opt-in items")
]

for idx, (m_name, m_val, m_desc) in enumerate(metrics, 5):
    ws2[f'B{idx}'] = m_name
    ws2[f'B{idx}'].font = Font(name="Calibri", size=11, bold=True)
    ws2[f'C{idx}'] = m_val
    ws2[f'C{idx}'].font = Font(name="Calibri", size=11, bold=True, color="1F497D")
    ws2[f'C{idx}'].alignment = Alignment(horizontal="center")
    ws2[f'D{idx}'] = m_desc
    ws2[f'D{idx}'].font = Font(name="Calibri", size=10)
    for c in ['B', 'C', 'D']:
        ws2[f'{c}{idx}'].border = thin_border

ws2.column_dimensions['B'].width = 45
ws2.column_dimensions['C'].width = 15
ws2.column_dimensions['D'].width = 65

wb.save(excel_path)
print(f"Saved styled Excel workbook: {excel_path}")
