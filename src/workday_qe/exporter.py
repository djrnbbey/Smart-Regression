"""
Enterprise Deliverable Exporter with CWE-1236 Formula Sanitization & SHA-256 Checksums
Zero-DLL Pure-Python & Openpyxl Implementation (Compatible with Windows Application Control)
"""
import csv
import shutil
from pathlib import Path
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.workday_qe.config import (
    OUTPUT_EXCEL,
    OUTPUT_CSV,
    ROOT_OUTPUT_EXCEL,
    ROOT_OUTPUT_CSV,
    GAP_IDENTIFIER
)
from src.workday_qe.guardrails import (
    sanitize_for_spreadsheet,
    compute_integrity_hash
)
from src.workday_qe.reporter import compute_executive_metrics

def export_to_csv(records: List[Dict[str, Any]], target_path: Path = OUTPUT_CSV) -> str:
    """
    Exports regression records to CSV with formula injection sanitization.
    Zero-DLL implementation using standard library csv module.
    """
    if not records:
        return ""
        
    safe_records = []
    for r in records:
        safe_r = {k: sanitize_for_spreadsheet(v) for k, v in r.items()}
        safe_records.append(safe_r)
        
    p = Path(target_path)
    p.parent.mkdir(parents=True, exist_ok=True)
    
    headers = list(safe_records[0].keys())
    with open(p, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(safe_records)
    
    # Also sync to root
    if p != ROOT_OUTPUT_CSV:
        shutil.copyfile(p, ROOT_OUTPUT_CSV)
        
    return str(p)

def export_to_excel(
    records: List[Dict[str, Any]],
    catalog_jobs: List[Dict[str, Any]],
    target_path: Path = OUTPUT_EXCEL
) -> Path:
    """
    Exports regression records to a beautifully styled Excel workbook with:
    - CWE-1236 formula sanitization
    - Priority color fills (High=Soft Red, Med=Soft Amber, Low=Soft Green)
    - Auto-filtering and frozen panes
    - Executive Summary Sheet with SHA-256 audit fingerprint
    """
    if not records:
        return Path(target_path)
        
    safe_records = []
    for r in records:
        safe_r = {k: sanitize_for_spreadsheet(v) for k, v in r.items()}
        safe_records.append(safe_r)
        
    wb = openpyxl.Workbook()
    
    # ----------------------------------------------------
    # Sheet 1: Targeted Regression Plan
    # ----------------------------------------------------
    ws1 = wb.active
    ws1.title = "Targeted Regression Plan"
    
    headers = list(safe_records[0].keys())
    ws1.append(headers)
    
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    for col_idx, h in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
    p_fills = {
        "High": PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
        "Med": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
        "Low": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    }
    p_fonts = {
        "High": Font(name="Calibri", size=10, bold=True, color="C00000"),
        "Med": Font(name="Calibri", size=10, bold=True, color="B28900"),
        "Low": Font(name="Calibri", size=10, bold=True, color="375623")
    }
    gap_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    gap_font = Font(name="Calibri", size=10, italic=True, color="7F7F7F")
    
    for row_idx, rec in enumerate(safe_records, 2):
        for col_idx, h in enumerate(headers, 1):
            val = rec[h]
            cell = ws1.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            
            if h == "Selection Priority (High/Med/Low)":
                p = str(val).strip()
                if p in p_fills:
                    cell.fill = p_fills[p]
                    cell.font = p_fonts[p]
                cell.alignment = Alignment(horizontal="center", vertical="top")
                
            if h == "Targeted E2E Job ID & Name" and "Gap Identified" in str(val):
                cell.fill = gap_fill
                cell.font = gap_font
                
    col_widths = {
        1: 38, 2: 18, 3: 32, 4: 38, 5: 18,
        6: 55, 7: 22, 8: 16, 9: 14, 10: 30
    }
    for col_idx, width in col_widths.items():
        ws1.column_dimensions[get_column_letter(col_idx)].width = width
        
    ws1.freeze_panes = "A2"
    ws1.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(safe_records)+1}"
    
    # ----------------------------------------------------
    # Sheet 2: Executive Summary & Audit Checksum
    # ----------------------------------------------------
    ws2 = wb.create_sheet(title="Executive Summary")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2['B2'] = "WORKDAY RELEASE INTELLIGENCE & QE REGRESSION PLAN"
    ws2['B2'].font = Font(name="Calibri", size=16, bold=True, color="1F497D")
    
    metrics = compute_executive_metrics(records, catalog_jobs)
    integrity_hash = compute_integrity_hash(records)
    
    headers_s2 = ["Metric", "Value", "Context / Description"]
    for idx, h in enumerate(headers_s2, 2):
        cell = ws2.cell(row=4, column=idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    metric_rows = [
        ("Total Release Items Analyzed", metrics['total_release_items'], "Total What's New items processed from release notes"),
        ("Release Items with Direct / Downstream E2E Test Impact", metrics['mapped_items_count'], "Features mapped to existing regression suite"),
        ("Identified Coverage Gaps (Missing E2E Tests)", metrics['gap_items_count'], "Release features lacking test coverage in catalog"),
        ("Master E2E Job Catalog Size", metrics['master_catalog_total'], "Total test inventory available in Master Repository"),
        ("Targeted E2E Regression Jobs Selected", metrics['unique_selected_jobs_count'], "Unique tests scheduled for this release cycle"),
        ("Regression Suite Reduction Percentage", f"{metrics['suite_reduction_pct']:.1f}%", metrics['bounds_validation_message']),
        ("High Priority Selected Items", metrics['high_priority_count'], "Core business transactions and regulatory changes"),
        ("Medium Priority Selected Items", metrics['med_priority_count'], "Configuration-dependent UI and calculation updates"),
        ("Low Priority Selected Items", metrics['low_priority_count'], "Cosmetic or low-blast-radius opt-in items"),
        ("SHA-256 Deliverable Integrity Checksum", integrity_hash[:32] + "...", f"Full SHA-256: {integrity_hash}")
    ]
    
    for idx, (m_name, m_val, m_desc) in enumerate(metric_rows, 5):
        ws2[f'B{idx}'] = m_name
        ws2[f'B{idx}'].font = Font(name="Calibri", size=11, bold=True)
        ws2[f'C{idx}'] = m_val
        ws2[f'C{idx}'].font = Font(name="Calibri", size=11, bold=True, color="1F497D")
        ws2[f'C{idx}'].alignment = Alignment(horizontal="center")
        ws2[f'D{idx}'] = m_desc
        ws2[f'D{idx}'].font = Font(name="Calibri", size=10)
        for c in ['B', 'C', 'D']:
            ws2[f'{c}{idx}'].border = thin_border
            
    ws2.column_dimensions['B'].width = 45
    ws2.column_dimensions['C'].width = 20
    ws2.column_dimensions['D'].width = 65
    
    p = Path(target_path)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    
    # Also sync to root
    if p != ROOT_OUTPUT_EXCEL:
        shutil.copyfile(p, ROOT_OUTPUT_EXCEL)
        
    return p
